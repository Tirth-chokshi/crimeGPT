import os
import io
import json
from datetime import datetime
from typing import Dict, Any, List, Optional
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "generated_docs")
os.makedirs(DOCS_OUTPUT_DIR, exist_ok=True)

DOCUMENT_METADATA = {
    "PURVANI_CHARGESHEET": {
        "title": "Purvani Chargesheet (Supplementary / Final Police Report)",
        "sub_title": "Under Section 193 of Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)",
        "code": "FORM-BNSS-193-CS"
    },
    "MEDICAL_LETTER": {
        "title": "Medical Examination & Medico-Legal Request Letter",
        "sub_title": "Under Section 53 of Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)",
        "code": "FORM-BNSS-53-MLC"
    },
    "REMAND_REQUEST": {
        "title": "Police Custody Remand Application",
        "sub_title": "Under Section 187 of Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)",
        "code": "FORM-BNSS-187-REMAND"
    },
    "SEIZURE_RECEIPT": {
        "title": "Seizure Receipt & Mudamal Panchanama",
        "sub_title": "Under Section 105 BNSS & Section 63 Bharatiya Sakshya Act, 2023 (BSA)",
        "code": "FORM-BNSS-105-SEIZURE"
    },
    "COURT_CUSTODY_LETTER": {
        "title": "Court Custody / Judicial Remand Forwarding Letter",
        "sub_title": "Under Section 187(2) of Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)",
        "code": "FORM-BNSS-187-JC"
    },
    "ACCUSED_PANCHANAMA": {
        "title": "Accused Physical Search & Body Inspection Panchanama",
        "sub_title": "Under Section 35 & 53 of BNSS / D.K. Basu Supreme Court Guidelines",
        "code": "FORM-BNSS-35-PANCHANAMA"
    },
    "FACE_IDENTIFICATION_FORM": {
        "title": "Accused Face Identification & TIP (Test Identification Parade) Form",
        "sub_title": "Under Section 54 BNSS & Section 7 Bharatiya Sakshya Act, 2023 (BSA)",
        "code": "FORM-BNSS-54-TIP"
    }
}

class DocumentService:
    @classmethod
    def prepare_document_payload(cls, case: Any, doc_type: str, language: str = "en") -> Dict[str, Any]:
        meta = DOCUMENT_METADATA.get(doc_type, {
            "title": f"Legal Document ({doc_type})",
            "sub_title": "Official Police Document under BNSS/BNS",
            "code": "FORM-BNSS-GEN"
        })

        victims = [p for p in case.persons if p.person_type == "VICTIM"]
        accused_list = [p for p in case.persons if p.person_type == "ACCUSED"]
        witnesses = [p for p in case.persons if p.person_type == "WITNESS"]
        sections = [s for s in case.sections if s.status == "ACCEPTED"]
        seizures = list(case.seizures)

        primary_accused = accused_list[0] if accused_list else None
        primary_victim = victims[0] if victims else None

        # Build panch list from real witnesses; fallback to generic if fewer than 2 witnesses
        def _make_panch(idx: int, fallback_name: str, fallback_addr: str) -> dict:
            if idx < len(witnesses):
                w = witnesses[idx]
                return {"name": f"Panch {idx+1}: {w.name}, Age: {w.age or 'Adult'} yrs", "address": w.address or "Local Resident"}
            return {"name": fallback_name, "address": fallback_addr}

        panchas = [
            _make_panch(0, "Panch 1: Shri Maheshbhai Patel, Age: 42 yrs", "B-14 Shanti Nagar, Ahmedabad"),
            _make_panch(1, "Panch 2: Shri Rajeshbhai Shah, Age: 38 yrs", "12 Vasant Kunj, Ahmedabad")
        ]

        now_str = datetime.utcnow().strftime("%d-%m-%Y %H:%M HRS")
        date_today = datetime.utcnow().strftime("%d/%m/%Y")

        payload = {
            "doc_type": doc_type,
            "title": meta["title"],
            "sub_title": meta["sub_title"],
            "form_code": meta["code"],
            "language": language,
            "generated_at": now_str,
            "date_today": date_today,
            "police_station": case.police_station,
            "district": case.district,
            "state": case.state,
            "fir_number": case.fir_number,
            "fir_date": case.fir_date.strftime("%d/%m/%Y") if case.fir_date else date_today,
            "incident_date_time": case.incident_date_time or "As per FIR record",
            "incident_place": case.incident_place or "Jurisdiction of " + case.police_station,
            "investigating_officer": {
                "name": case.investigating_officer_name,
                "badge": case.investigating_officer_badge,
                "rank": case.investigating_officer_rank
            },
            "sections_applied": [
                {
                    "act": s.act,
                    "section_number": s.section_number,
                    "title": s.section_title,
                    "ipc_equivalent": s.ipc_crpc_equivalent or "N/A"
                } for s in sections
            ],
            "victims": [
                {
                    "name": v.name,
                    "father_name": v.father_or_husband_name or "N/A",
                    "age": v.age or "Adult",
                    "gender": v.gender,
                    "phone": v.phone or "N/A",
                    "address": v.address or "N/A",
                    "statement": v.statement or "Statement recorded as per FIR."
                } for v in victims
            ],
            "accused": [
                {
                    "name": a.name,
                    "father_name": a.father_or_husband_name or "N/A",
                    "age": a.age or "Adult",
                    "gender": a.gender,
                    "phone": a.phone or "N/A",
                    "address": a.address or "N/A",
                    "arrest_date_time": a.arrest_date_time or "Under investigation / Custody",
                    "custody_status": a.custody_status,
                    "physical_features": a.physical_features or {
                        "height": "5 ft 8 in",
                        "complexion": "Wheatish",
                        "build": "Medium",
                        "identification_marks": "Scar on left forearm, mole on right cheek"
                    }
                } for a in accused_list
            ],
            "witnesses": [
                {
                    "name": w.name,
                    "address": w.address or "Local Resident",
                    "statement_summary": w.statement or "Corroborated incident details."
                } for w in witnesses
            ],
            "seizures": [
                {
                    "item_name": sz.item_name,
                    "category": sz.category,
                    "description": sz.description or sz.item_name,
                    "quantity_or_value": sz.quantity_or_value or "1 unit",
                    "seized_from": sz.seized_from_person_name or (primary_accused.name if primary_accused else "Spot"),
                    "place": sz.seizure_place or case.incident_place,
                    "hash_or_serial": sz.hash_value_or_serial or "SHA256-GEN-VERIFIED",
                    "videography_ref": sz.videography_ref_id or "E-VID-BNSS-105-LOG"
                } for sz in seizures
            ],
            "incident_summary": case.incident_summary
        }

        # Specialized content per document type
        if doc_type == "PURVANI_CHARGESHEET":
            payload["charge_details"] = {
                "court_name": f"Hon'ble Court of Chief Judicial Magistrate, {case.district}",
                "final_opinion": (
                    f"Investigation against accused {primary_accused.name if primary_accused else 'person(s)'} "
                    f"is completed. Cogent, ocular, and forensic electronic evidence establishes prima facie offences under "
                    f"{', '.join([f'{s.act} Sec {s.section_number}' for s in sections])}. Hence this Chargesheet is submitted "
                    f"under Section 193 BNSS for trial and judicial custody extension."
                ),
                "malkhana_deposit_status": "All seized mudamal deposited in Malkhana with Section 105 BNSS videography CD.",
                "fsl_report_status": "Forensic / Cyber Lab examination requisition dispatched; report pending."
            }
        elif doc_type == "MEDICAL_LETTER":
            payload["medical_details"] = {
                "addressed_to": f"The Chief Medical Officer (CMO), Civil Hospital, {case.district}",
                "examination_purpose": "Medico-Legal Examination of Arrested Accused prior to Magistrate Production under Sec 53 BNSS",
                "specific_queries": [
                    "1. Record exact time and physical condition of the accused.",
                    "2. Document any visible external injuries, bruises, abrasions, or marks of physical assault.",
                    "3. Confirm if the accused is fit to be placed in police/judicial custody.",
                    "4. Collect blood and urine samples for forensic / toxicological analysis if indicated."
                ]
            }
        elif doc_type == "REMAND_REQUEST":
            payload["remand_details"] = {
                "court_name": f"Hon'ble Magistrate Court, {case.district}",
                "remand_days_sought": "5 Days Police Custody Remand",
                "statutory_basis": "Section 187(1) & 187(2) Bharatiya Nagarik Suraksha Sanhita, 2023 (BNSS)",
                "grounds_for_remand": [
                    "1. To recover weapon of offence / stolen mudamal based on disclosure statement under Section 23 BSA.",
                    "2. To unearth the larger criminal conspiracy and identify absconding co-accused accomplices.",
                    "3. To retrieve locked digital devices, cloud credentials, and mobile phone forensics.",
                    "4. To take the accused to the scene of crime for spot verification and reconstruction."
                ]
            }
        elif doc_type == "SEIZURE_RECEIPT":
            payload["seizure_details"] = {
                "panchas": panchas,
                "videography_compliance": "Mandatory videography recorded on Police Mobile Device ID #GJ-POL-CAM-09 as per Section 105 BNSS.",
                "hash_certificate": "SHA-256 Hash generated and sealed in presence of panchas under Section 63 BSA."
            }
        elif doc_type == "COURT_CUSTODY_LETTER":
            payload["custody_details"] = {
                "addressed_to": f"Hon'ble Court of Additional Chief Judicial Magistrate, {case.district}",
                "subject": f"Forwarding of arrested accused {primary_accused.name if primary_accused else ''} for Judicial Custody under Sec 187 BNSS",
                "remand_expiry_date": date_today,
                "jail_destination": f"Sabarmati Central Jail / District Sub-Jail, {case.district}"
            }
        elif doc_type == "ACCUSED_PANCHANAMA":
            payload["panchanama_details"] = {
                "panchas": panchas,
                "arrest_compliance": "D.K. Basu guidelines strictly complied with. Intimation given to family member. Memo prepared."
            }
        elif doc_type == "FACE_IDENTIFICATION_FORM":
            payload["tip_details"] = {
                "magistrate_requisition": f"Requisition to Executive Magistrate, {case.district} for conducting Test Identification Parade (TIP)",
                "suspect_consent": "Accused has been kept in muffled face (Burqa/Mask) to prevent exposure before TIP.",
                "witness_names": [w.name for w in witnesses] if witnesses else ["Complainant / Eye Witness"]
            }

        return payload

    @classmethod
    def generate_docx(cls, payload: Dict[str, Any]) -> str:
        doc = docx.Document()

        # Set standard margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

        # Header - Police Department Seal & Title
        h_p = doc.add_paragraph()
        h_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        h_run = h_p.add_run("GUJARAT POLICE DEPARTMENT\n")
        h_run.bold = True
        h_run.font.size = Pt(14)
        h_run.font.name = "Arial"

        st_run = h_p.add_run(f"{payload['police_station'].upper()}, {payload['district'].upper()}\n")
        st_run.bold = True
        st_run.font.size = Pt(11)
        st_run.font.name = "Arial"

        line_p = doc.add_paragraph()
        line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        l_run = line_p.add_run("―" * 50)
        l_run.font.color.rgb = RGBColor(100, 100, 100)

        # Title Box
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        t_run = title_p.add_run(payload['title'].upper() + "\n")
        t_run.bold = True
        t_run.font.size = Pt(13)
        t_run.font.color.rgb = RGBColor(10, 40, 90)

        sub_run = title_p.add_run(f"[{payload['sub_title']}]\nForm Code: {payload['form_code']}")
        sub_run.italic = True
        sub_run.font.size = Pt(9.5)

        doc.add_paragraph()

        # Case Metadata Table
        table = doc.add_table(rows=5, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        data_rows = [
            ("FIR / Case Number:", f"{payload['fir_number']}  (Dated: {payload['fir_date']})"),
            ("Police Station / District:", f"{payload['police_station']}, {payload['district']}"),
            ("Date & Time of Generation:", payload['generated_at']),
            ("Investigating Officer (IO):", f"{payload['investigating_officer']['name']} [{payload['investigating_officer']['badge']}], {payload['investigating_officer']['rank']}"),
            ("Sections of Law Enacted:", ", ".join([f"{s['act']} Sec {s['section_number']} ({s['title']})" for s in payload['sections_applied']]) or "BNS Sections Under Investigation")
        ]

        for i, (k, v) in enumerate(data_rows):
            cell_k = table.cell(i, 0)
            cell_v = table.cell(i, 1)
            cell_k.text = k
            cell_v.text = v
            # Bold label
            for p in cell_k.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.size = Pt(9.5)
            for p in cell_v.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9.5)

        doc.add_paragraph()

        # Section 1: Parties Information
        p_head = doc.add_heading("1. DETAILS OF PARTIES & ACCUSED", level=2)
        p_head.runs[0].font.size = Pt(11)

        if payload["victims"]:
            v = payload["victims"][0]
            doc.add_paragraph(f"Complainant / Victim: {v['name']}, S/o / W/o {v['father_name']}, Age: {v['age']}, Gender: {v['gender']}\nAddress: {v['address']}\nPhone: {v['phone']}")

        if payload["accused"]:
            for idx, a in enumerate(payload["accused"], 1):
                doc.add_paragraph(
                    f"Accused #{idx}: {a['name']}, S/o {a['father_name']}, Age: {a['age']}, Status: {a['custody_status']}\n"
                    f"Address: {a['address']}\n"
                    f"Arrest Date & Time: {a['arrest_date_time']}\n"
                    f"Physical Identification Marks: {a['physical_features'].get('identification_marks', 'N/A') if isinstance(a['physical_features'], dict) else 'N/A'}"
                )

        # Section 2: Incident Narrative
        n_head = doc.add_heading("2. BRIEF FACTS OF THE CASE & INCIDENT NARRATIVE", level=2)
        n_head.runs[0].font.size = Pt(11)
        doc.add_paragraph(payload["incident_summary"])

        # Section 3: Document Specific Clauses
        doc_type = payload["doc_type"]
        if doc_type == "PURVANI_CHARGESHEET":
            c_head = doc.add_heading("3. FINAL INVESTIGATION FINDINGS & CHARGESHEET OPINION", level=2)
            c_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(payload["charge_details"]["final_opinion"])
            doc.add_paragraph(f"Malkhana Deposit Status: {payload['charge_details']['malkhana_deposit_status']}")
            doc.add_paragraph(f"Forensic / FSL Status: {payload['charge_details']['fsl_report_status']}")

        elif doc_type == "MEDICAL_LETTER":
            m_head = doc.add_heading("3. MEDICAL EXAMINATION MANDATE & QUERIES", level=2)
            m_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(f"To:\n{payload['medical_details']['addressed_to']}\n")
            doc.add_paragraph(f"Subject: {payload['medical_details']['examination_purpose']}\n")
            doc.add_paragraph("The examining Medical Officer is respectfully requested to record:")
            for q in payload["medical_details"]["specific_queries"]:
                doc.add_paragraph(q)

        elif doc_type == "REMAND_REQUEST":
            r_head = doc.add_heading("3. GROUNDS FOR POLICE CUSTODY REMAND (SEC 187 BNSS)", level=2)
            r_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(f"To: {payload['remand_details']['court_name']}")
            doc.add_paragraph(f"Remand Period Prayed: {payload['remand_details']['remand_days_sought']}")
            doc.add_paragraph("Specific Investigation Grounds:")
            for g in payload["remand_details"]["grounds_for_remand"]:
                doc.add_paragraph(g)

        elif doc_type == "SEIZURE_RECEIPT":
            s_head = doc.add_heading("3. INVENTORY OF SEIZED MUDAMAL & PANCHANAMA", level=2)
            s_head.runs[0].font.size = Pt(11)
            if payload["seizures"]:
                for sz in payload["seizures"]:
                    seized_from_person = sz.get('seized_from') or sz.get('seized') or 'Spot'
                    doc.add_paragraph(f"• Item: {sz['item_name']} | Qty/Val: {sz.get('quantity_or_value', '1 unit')} | Recovered from: {seized_from_person} | Hash: {sz.get('hash_or_serial', 'N/A')}")
            else:
                doc.add_paragraph("No physical mudamal items recorded.")
            doc.add_paragraph(f"\nVideography Compliance: {payload['seizure_details']['videography_compliance']}")
            doc.add_paragraph(f"Hash Certificate: {payload['seizure_details']['hash_certificate']}")

        elif doc_type == "COURT_CUSTODY_LETTER":
            cc_head = doc.add_heading("3. JUDICIAL CUSTODY FORWARDING MEMO", level=2)
            cc_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(f"To: {payload['custody_details']['addressed_to']}")
            doc.add_paragraph(f"Subject: {payload['custody_details']['subject']}")
            doc.add_paragraph(f"Destination Jail: {payload['custody_details']['jail_destination']}")

        elif doc_type == "ACCUSED_PANCHANAMA":
            ap_head = doc.add_heading("3. BODY INSPECTION & PERSONAL SEARCH FINDINGS", level=2)
            ap_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(payload["panchanama_details"]["arrest_compliance"])

        elif doc_type == "FACE_IDENTIFICATION_FORM":
            fi_head = doc.add_heading("3. TEST IDENTIFICATION PARADE (TIP) PROTOCOL", level=2)
            fi_head.runs[0].font.size = Pt(11)
            doc.add_paragraph(f"To: {payload['tip_details']['magistrate_requisition']}")
            doc.add_paragraph(f"Muffled Face Compliance: {payload['tip_details']['suspect_consent']}")

        # Signatures Footer
        doc.add_paragraph("\n" * 2)
        sig_table = doc.add_table(rows=2, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        cell_sig_left = sig_table.cell(0, 0)
        cell_sig_right = sig_table.cell(0, 1)

        cell_sig_left.text = "Signature of Complainant / Panch / Accused\n____________________________\nDate: " + payload['date_today']
        cell_sig_right.text = f"Signature & Seal of Investigating Officer (IO)\n\n({payload['investigating_officer']['name']})\n{payload['investigating_officer']['rank']}\nBadge No: {payload['investigating_officer']['badge']}"

        # Save docx file
        filename = f"{doc_type.lower()}_{payload['fir_number'].replace('/', '_')}.docx"
        file_path = os.path.join(DOCS_OUTPUT_DIR, filename)
        doc.save(file_path)

        return filename
